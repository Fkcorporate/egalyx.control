# services/veille_export_service.py
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


class VeilleExportService:
    """Service d'export pour les veilles réglementaires"""
    
    @staticmethod
    def export_excel(veilles, titre='Veilles réglementaires'):
        """Export Excel des veilles"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Veilles"
        
        # Styles
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill('solid', fgColor='1A2A5E')
        border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )
        
        # Titre du document
        ws.merge_cells('A1:K1')
        title_cell = ws['A1']
        title_cell.value = titre
        title_cell.font = Font(bold=True, size=14, color='1A2A5E')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # Date
        ws.merge_cells('A2:K2')
        date_cell = ws['A2']
        date_cell.value = f"Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
        date_cell.font = Font(italic=True, size=10, color='666666')
        date_cell.alignment = Alignment(horizontal='center')
        
        # En-têtes (ligne 4)
        headers = [
            'Référence', 'Titre', 'Type', 'Organisme', 'Statut',
            'Impact', 'Date publication', 'Date application',
            'Actions actives', 'Actions terminées', 'Documents'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Données
        for row, veille in enumerate(veilles, 5):
            actions_actives = veille.actions.filter_by(is_archived=False).count()
            actions_terminees = veille.actions.filter_by(
                is_archived=False, statut='termine'
            ).count()
            
            data = [
                veille.reference or '',
                veille.titre,
                veille.type_reglementation or '',
                veille.organisme_emetteur or '',
                veille.statut or '',
                veille.impact_estime or '',
                veille.date_publication.strftime('%d/%m/%Y') if veille.date_publication else '',
                veille.date_application.strftime('%d/%m/%Y') if veille.date_application else '',
                actions_actives,
                actions_terminees,
                veille.documents.count()
            ]
            
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = border
                if col == 1:
                    cell.font = Font(bold=True)
        
        # Largeurs
        widths = [15, 40, 15, 20, 12, 12, 14, 14, 12, 12, 12]
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=4, column=col).column_letter].width = width
        
        # Sauvegarder en mémoire
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def export_pdf(veilles, titre='Rapport Veille Réglementaire'):
        """Export PDF des veilles"""
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4)
        elements = []
        
        styles = getSampleStyleSheet()
        
        # Titre
        title_style = styles['Title']
        title_style.textColor = colors.HexColor('#1A2A5E')
        elements.append(Paragraph(titre, title_style))
        elements.append(Spacer(1, 10))
        
        # Date
        date_style = styles['Normal']
        date_style.fontSize = 9
        date_style.textColor = colors.grey
        elements.append(Paragraph(
            f"Exporté le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
            date_style
        ))
        elements.append(Spacer(1, 20))
        
        # Tableau
        data = [['Référence', 'Titre', 'Statut', 'Impact', 'Date app.']]
        
        for v in veilles:
            data.append([
                v.reference or '-',
                v.titre[:50] + ('...' if len(v.titre) > 50 else ''),
                v.statut or '-',
                v.impact_estime or '-',
                v.date_application.strftime('%d/%m/%Y') if v.date_application else '-'
            ])
        
        table = Table(data, colWidths=[80, 220, 70, 60, 80])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A2A5E')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
        ]))
        elements.append(table)
        
        doc.build(elements)
        output.seek(0)
        return output